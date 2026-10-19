from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR

BASE=str(Path(__file__).resolve().parents[1])
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
BG=RGBColor(12,25,42); PANEL=RGBColor(24,43,65); WHITE=RGBColor(240,247,253); DIM=RGBColor(166,191,210); TEAL=RGBColor(33,211,189); YELLOW=RGBColor(243,195,92)


def rect(s,x,y,w,h,fill=PANEL,radius=True):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background();return sh

def text(s,x,y,w,h,msg,size=18,color=WHITE,bold=False,align=None):
 tb=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=tb.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=Inches(.05);tf.margin_top=Inches(.04);tf.margin_bottom=Inches(0)
 for i,line in enumerate(str(msg).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
  if align is not None:p.alignment=align
 return tb

def slide(title,lead,index):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=BG
 rect(s,.55,.35,.11,.53,TEAL,False)
 text(s,.83,.26,11,.65,title,30,WHITE,True)
 text(s,.83,1.02,11.5,.52,lead,14,DIM)
 rect(s,.72,7.03,11.9,.016,TEAL,False)
 text(s,.78,7.08,8,.25,'SUNSET WORLD  |  PROPUESTA TÉCNICA  |  DATOS SINTÉTICOS',10,DIM)
 text(s,11.8,7.08,.8,.25,f'{index:02d}/05',10,TEAL,True,PP_ALIGN.RIGHT)
 return s

def card(s,x,y,w,h,header,body,accent=TEAL):
 rect(s,x,y,w,h)
 rect(s,x+.20,y+.28,.055,.45,accent,False)
 text(s,x+.38,y+.2,w-.56,.57,header,15 if w<3 else 19,WHITE,True)
 text(s,x+.31,y+.87,w-.52,h-1,body,11 if w<3 else 14,DIM)

s=slide('Una sola plataforma de datos','Arquitectura modular, auditable y reproducible — opción A',1)
stages=[('FUENTES','Voltux · Share CRM\nTMS · FX · Call Center'),('LAKE','Python / Oracle 11g\nParquet · S3/MinIO'),('WAREHOUSE','PostgreSQL / Redshift\nEstrella · SCD Tipo 2'),('SEMÁNTICA','dbt + reglas de calidad\nLinaje · monitoreo'),('IA','Mantis · pgvector\nRAG · búsqueda semántica')]
for i,(h,b) in enumerate(stages):
 x=.65+i*2.55;card(s,x,2.05,2.35,2.27,h,b)
 if i<4:text(s,x+2.36,2.81,.20,.4,'›',23,TEAL,True)
rect(s,.8,5.15,11.7,.85)
text(s,1.11,5.33,11.1,.43,'Operación: Airflow 02:00  •  Kafka + Spark en tiempo real  •  calidad bloqueante antes de publicar',17,WHITE,True)
text(s,.84,6.25,11.8,.45,'Propuesta de implementación. No se han probado en este entorno los servicios empresariales.',12,YELLOW)

s=slide('Volúmenes: referencia vs evidencia','Separar la magnitud del reto de lo que ya está ejecutado',2)
for i,(n,v,sub) in enumerate([('VOLTUX','4.5 M','Registros del sistema (referencia)'),('TMS PAGOS','6.9 M','Transacciones (referencia)'),('TMS COMENTARIOS','10 M','Textos (referencia)')]):
 x=.78+i*4.26;rect(s,x,1.86,3.89,2.15)
 text(s,x+.26,2.1,3.4,.35,n,14,DIM,True)
 text(s,x+.25,2.47,3.4,.71,v,34,TEAL,True)
 text(s,x+.26,3.41,3.4,.4,sub,12,DIM)
rect(s,.78,4.38,11.85,1.54)
text(s,1.04,4.63,10.75,.45,'MUESTRA REALMENTE VALIDADA',19,WHITE,True)
text(s,1.04,5.16,11,.45,'1,000 leads  ·  1,462 cotizaciones FX  ·  600 pagos  ·  1,200 comentarios',18,TEAL)
text(s,.85,6.2,11.6,.4,'La ejecución masiva está parametrizada, pero NO se ha ejecutado: no se muestran benchmarks ficticios.',12,YELLOW)

s=slide('Calidad y modelo dimensional','Probar relaciones y consistencia antes de escalar el volumen',3)
card(s,.78,1.88,5.65,3.62,'Modelo dimensional','fact_ventas · fact_fx\nfact_interacciones · fact_tours\n\nSCD2: cliente y sala\nClaves sustitutas y correspondencias entre sistemas')
card(s,6.74,1.88,5.85,3.62,'Evidencia comprobada','31 pruebas Python: PASS\n14 controles de calidad: PASS\n4 relaciones conciliadas: PASS\n\nContratos/ventas, pagos/contratos,\ncomentarios/tours y ventas/FX',YELLOW)
text(s,.91,6.08,11.4,.64,'SQLite valida la lógica de la muestra. PostgreSQL y dbt siguen pendientes de ejecución en Docker.',13,DIM)

s=slide('Del Call Center a Mantis','Procesamiento near-real-time y búsqueda semántica con controles',4)
card(s,.78,1.86,5.71,3.65,'Kafka + Spark','transcripciones_raw\n   ↓ regex / reglas aprobadas\ntranscripciones_procesadas + alertas\n   ↓ UPSERT por event_id\nPostgreSQL (replay idempotente)')
card(s,6.72,1.86,5.86,3.65,'Embeddings + RAG','Comentarios/transcripciones a Parquet\n   ↓ privacidad + chunking\nEmbeddings incrementales (hash)\n   ↓ pgvector\nMantis: recuperación con fuentes')
text(s,.85,6.10,11.8,.58,'Limitación: código implementado, integración Kafka/Spark/pgvector y API Mantis aún no verificadas.',13,YELLOW)

s=slide('Viabilidad, riesgos y próximos hitos','Propuesta orientada a valor para FX, Client Journey y consolidación',5)
card(s,.78,1.83,3.81,3.7,'Beneficio esperado','FX centralizado y conciliable\nVisión de cliente y contratos\nMenos inconsistencias\nDatos gobernados para Mantis')
card(s,4.75,1.83,3.82,3.7,'Riesgos y costos','Acceso Oracle 11g / latencia\nPII, consentimiento y ACL\nRecursos Spark / Kafka\nStorage, cómputo, operación\nNo hay ROI estimado sin piloto',YELLOW)
card(s,8.73,1.83,3.82,3.7,'Hitos de aceptación','1. Probar Parquet y DW\n2. Ejecutar Airflow + dbt\n3. Validar streaming y RAG\n4. Benchmark volumetría\n5. Grabar demo real')
text(s,.87,6.12,11.8,.5,'Decisiones de negocio pendientes: definición VPG y conversión, reglas de transcripción y API real de Mantis.',12,DIM)

prs.save(BASE+'/deliverables/Sunset_World_Presentacion_Ejecutiva.pptx')
print('slides generated',len(prs.slides))
